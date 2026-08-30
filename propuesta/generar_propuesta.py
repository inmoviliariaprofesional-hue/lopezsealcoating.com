#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generador del entregable Word: Propuesta de Presencia Digital para Lopez Sealcoating LLC.
Preparado por Aimarktech. Ejecutar: python3 generar_propuesta.py
Salida: Propuesta_Presencia_Digital_Lopez_Sealcoating.docx
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------- Paleta de marca ----------------
ASFALTO = RGBColor(0x1A, 0x1A, 0x1A)   # negro asfalto
AMARILLO = RGBColor(0xF5, 0xC4, 0x00)  # amarillo señalización
GRIS = RGBColor(0x55, 0x55, 0x55)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
AMARILLO_HEX = "F5C400"
ASFALTO_HEX = "1A1A1A"
GRIS_CLARO_HEX = "EDEDED"

doc = Document()

# Márgenes
for section in doc.sections:
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)

# Fuente base
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = ASFALTO


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, color=ASFALTO, size=10.5, align=None, white=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = BLANCO if white else color
    run.font.name = "Calibri"
    return p


def add_bar_heading(text, fill_hex=ASFALTO_HEX, text_white=True):
    """Encabezado de sección como una barra de color de ancho completo."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    shade_cell(cell, fill_hex)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("  " + text)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = BLANCO if text_white else ASFALTO
    run.font.name = "Calibri"
    doc.add_paragraph()
    return tbl


def add_sub(text, color=ASFALTO):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = color
    return p


def add_body(text, size=11, space_after=6, italic=False, bold=False, color=ASFALTO):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.italic = italic
    run.bold = bold
    run.font.color.rgb = color
    return p


def add_bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        r.font.size = Pt(11)
    r2 = p.add_run(text)
    r2.font.size = Pt(11)
    return p


# ============================================================
# PORTADA
# ============================================================
# Marca Aimarktech
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("AIMARKTECH")
r.bold = True
r.font.size = Pt(22)
r.font.color.rgb = ASFALTO
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run("Marketing · Tecnología · Presencia Digital")
r2.font.size = Pt(11)
r2.italic = True
r2.font.color.rgb = GRIS

doc.add_paragraph()
doc.add_paragraph()

# Banda de título
tbl = doc.add_table(rows=1, cols=1)
cell = tbl.rows[0].cells[0]
shade_cell(cell, ASFALTO_HEX)
pc = cell.paragraphs[0]
pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
pc.paragraph_format.space_before = Pt(14)
pc.paragraph_format.space_after = Pt(6)
r = pc.add_run("PROPUESTA DE PRESENCIA DIGITAL")
r.bold = True
r.font.size = Pt(20)
r.font.color.rgb = BLANCO
pc2 = cell.add_paragraph()
pc2.alignment = WD_ALIGN_PARAGRAPH.CENTER
pc2.paragraph_format.space_after = Pt(14)
r = pc2.add_run("Sitio web + Google + Reseñas")
r.font.size = Pt(13)
r.font.color.rgb = AMARILLO
r.bold = True

doc.add_paragraph()

# Cliente
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Preparada para:")
r.font.size = Pt(11)
r.font.color.rgb = GRIS
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("LOPEZ SEALCOATING LLC")
r.bold = True
r.font.size = Pt(16)
r.font.color.rgb = ASFALTO
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Atención: Demetrio López  ·  Chicagoland, Illinois (EE. UU.)")
r.font.size = Pt(11)
r.font.color.rgb = GRIS

doc.add_paragraph()
doc.add_paragraph()

# Datos de la propuesta (tabla)
data = [
    ("Elaborada por", "Antonio Aguilar — CEO, Aimarktech"),
    ("Sitio web", "www.soyaimarktech.com"),
    ("Fecha", "Agosto de 2026"),
    ("Vigencia de la propuesta", "30 días naturales"),
]
meta = doc.add_table(rows=len(data), cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, (k, v) in enumerate(data):
    set_cell_text(meta.rows[i].cells[0], k, bold=True, size=10.5)
    set_cell_text(meta.rows[i].cells[1], v, size=10.5)
    shade_cell(meta.rows[i].cells[0], GRIS_CLARO_HEX)

doc.add_page_break()

# ============================================================
# 1. CARTA / INTRODUCCIÓN
# ============================================================
add_bar_heading("1. Introducción")
add_body(
    "Estimado Demetrio:", bold=True, space_after=4)
add_body(
    "Gracias por el tiempo que nos dedicaste para platicar sobre tu negocio. En esta propuesta "
    "resumimos lo que entendimos de Lopez Sealcoating y te presentamos un plan claro, por etapas, "
    "para que más personas te encuentren, confíen en ti y te contacten para pedir cotización."
)
add_body(
    "El objetivo no es solamente entregarte \u201cuna página web\u201d, sino construir una presencia "
    "digital que realmente te genere trabajo, respetando que trabajas por tu cuenta y que prefieres "
    "empezar con una inversión razonable. Por eso la propuesta está dividida en fases: puedes avanzar "
    "a tu ritmo, viendo resultados en cada paso."
)

# ============================================================
# 2. DIAGNÓSTICO
# ============================================================
add_bar_heading("2. Diagnóstico: dónde estás hoy")
add_body(
    "Durante nuestra conversación identificamos algo muy importante. Tú mismo nos comentaste que, "
    "cuando buscas un servicio, entras a Google y eliges a la empresa que tiene más estrellas, porque "
    "asumes que esa es la que da mejor servicio."
)
add_sub("El punto clave:")
add_body(
    "Tus clientes hacen exactamente lo mismo… pero hoy Lopez Sealcoating es prácticamente invisible en "
    "Google. No cuentas con un Perfil de Empresa en Google (Google Business Profile) ni con reseñas. "
    "Es decir, no apareces justo en el lugar donde la gente decide a quién contratar.",
    bold=True,
)
add_body("Situación actual que detectamos:")
add_bullet("Consigues clientes por recomendación y por volanteo casa por casa.", bold_prefix="Cómo llegan hoy: ")
add_bullet("Facebook e Instagram (con pocos resultados). Sin Google Business Profile.", bold_prefix="Redes: ")
add_bullet("Ninguna reseña pública, y no les pides a tus clientes que te califiquen.", bold_prefix="Reseñas: ")
add_bullet("No tienes dominio propio ni correo profesional (usas Gmail).", bold_prefix="Sitio/dominio: ")
add_bullet("Logo tipo mascota tomado de internet y modificado (conviene rehacerlo original).", bold_prefix="Logo: ")
add_bullet("Público bilingüe (inglés y español) residencial y comercial, en Chicagoland (radio 30\u201340 millas).", bold_prefix="Mercado: ")
add_body(
    "En pocas palabras: tienes un buen servicio y clientes satisfechos, pero tu negocio no se ve donde "
    "más importa. Esa es justamente la mayor oportunidad.",
    italic=True,
)

# ============================================================
# 3. OBJETIVOS
# ============================================================
add_bar_heading("3. Lo que buscamos lograr")
add_bullet("Que te encuentren en Google cuando alguien busque \u201csealcoating\u201d en tu zona.")
add_bullet("Generar reseñas reales que te den credibilidad y estrellas.")
add_bullet("Convertir esas visitas en llamadas y solicitudes de cotización.")
add_bullet("Dar imagen profesional y bilingüe (inglés/español), residencial y comercial.")
add_bullet("Abrir la puerta a clientes comerciales con tu servicio de pintura de líneas y cajones de estacionamiento.")

# ============================================================
# 4. SOLUCIÓN POR FASES
# ============================================================
add_bar_heading("4. Solución propuesta (por fases)")
add_body(
    "Ordenamos el trabajo de lo que más impacto genera al menor costo, hacia lo opcional de crecimiento. "
    "Puedes contratar fase por fase.",
    italic=True,
)

add_sub("FASE 0 — Cimientos (lo más importante y económico)", color=ASFALTO)
add_bullet("Logo profesional ORIGINAL, inspirado en el que ya tienes (mascota + casa de fondo, negro asfalto y amarillo señalización), pero rediseñado desde cero para que sea 100% tuyo y sin riesgos.")
add_bullet("Registro de dominio propio (ejemplo: lopezsealcoating.com).")
add_bullet("Correo profesional (ejemplo: info@lopezsealcoating.com), configurado sin costo mensual, reenviando a tu correo actual.")
add_bullet("Creación y optimización de tu Perfil de Empresa en Google (Google Business Profile): zona de servicio, servicios, fotos, horario de temporada y botón de llamada.")
add_bullet("Sistema simple para pedir reseñas: enlace directo + tarjeta/QR + guion para que tus clientes te califiquen fácil.")
add_body("Resultado: empiezas a aparecer en Google Maps y a acumular estrellas. Esto solo ya cambia el juego.", italic=True, bold=True)
add_body(
    "Cómo implementamos el Perfil de Google (paso a paso): 1) creamos y reclamamos el perfil; "
    "2) lo verificamos con Google (video/teléfono/tarjeta postal); 3) elegimos la categoría correcta "
    "(contratista de pavimento / sealcoating) y definimos tu zona de servicio por ciudades; "
    "4) cargamos servicios, descripción bilingüe, fotos, logo y horario de temporada; "
    "5) conectamos tu teléfono y el sitio web; 6) instalamos el sistema de reseñas y publicamos las "
    "primeras. Todo el perfil es gratuito en Google; se cobra el trabajo de armarlo y optimizarlo.",
    size=10, color=GRIS,
)

add_sub("FASE 1 — Sitio web de conversión (bilingüe)", color=ASFALTO)
add_bullet("Sitio web profesional en inglés y español (una sola página bien estructurada, rápida y para celular).")
add_bullet("Botón para llamar con un toque desde el teléfono y botón de WhatsApp.")
add_bullet("Formulario para solicitar cotización gratis (\u201cFree Estimate\u201d).")
add_bullet("Galería con tus fotos y videos reales de trabajos terminados.")
add_bullet("Secciones de servicios: sealcoating y pintura de líneas/cajones/cebras (residencial y comercial).")
add_bullet("Zonas que cubres, sellos de confianza (\u201cLicensed & Insured\u201d solo si aplica) y \u201cHablamos Español\u201d.")
add_bullet("Enlace directo a tu Perfil de Google y a tus reseñas.")
add_body("Resultado: cuando alguien te encuentra, tiene todo para contactarte y pedir cotización en segundos.", italic=True, bold=True)

add_sub("FASE 2 — Crecimiento (opcional, por temporada)", color=ASFALTO)
add_bullet("Campañas de anuncios en Google (Búsqueda / Local Services Ads) para la temporada de mayo a octubre.")
add_bullet("Publicidad segmentada en Facebook/Instagram por zona.")
add_bullet("Seguimiento de resultados: cuántas llamadas y cotizaciones llegan.")
add_body("Resultado: aceleras la llegada de clientes cuando ya tienes la base lista. Se recomienda solo después de la Fase 0 y 1.", italic=True, bold=True)

# ============================================================
# 5. INVERSIÓN
# ============================================================
add_bar_heading("5. Inversión (precios de referencia)")
add_body(
    "Para que quede totalmente claro qué pagas y a quién, separamos la inversión en dos partes: "
    "(A) los servicios profesionales que realiza Aimarktech, y (B) los costos de terceros que se "
    "contratan a nombre de tu negocio (dominio, correo, hospedaje y, si decides, anuncios)."
)


def build_price_table(rows, col2="Inversión (USD)*"):
    tb = doc.add_table(rows=1, cols=3)
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    tb.style = "Table Grid"
    h = tb.rows[0].cells
    set_cell_text(h[0], "Concepto", bold=True, white=True, size=11)
    set_cell_text(h[1], "Tipo", bold=True, white=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(h[2], col2, bold=True, white=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    for c in h:
        shade_cell(c, ASFALTO_HEX)
    for concepto, tipo, precio in rows:
        r = tb.add_row().cells
        set_cell_text(r[0], concepto, size=10.5)
        set_cell_text(r[1], tipo, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(r[2], precio, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    return tb


# --- Tabla A: servicios Aimarktech ---
add_sub("A) Servicios profesionales (los realiza Aimarktech)")
build_price_table([
    ("Diseño de logo original", "Único", "$120 – $180"),
    ("Perfil de Empresa en Google + optimización", "Único", "$120 – $200"),
    ("Sistema de reseñas (enlace + QR + guion)", "Único", "$60 – $100"),
    ("Sitio web bilingüe (EN/ES) de conversión", "Único", "$450 – $800"),
    ("Gestión de anuncios (opcional, Fase 2)", "Mensual", "$200 – $400"),
])
add_body(
    "Son honorarios por el trabajo profesional. El Perfil de Empresa en Google es gratuito por parte "
    "de Google; lo que se cobra es la creación, verificación y optimización del perfil.",
    size=10, color=GRIS,
)

# --- Tabla B: costos de terceros ---
add_sub("B) Costos de terceros (a nombre de tu negocio)")
build_price_table([
    ("Dominio .com (registrador a costo, ej. Cloudflare)", "Anual", "$10 – $15"),
    ("Correo profesional (opción sin costo mensual)", "Anual", "$0"),
    ("Correo con buzón completo (Google Workspace, opcional)", "Mensual", "≈ $7 / mes"),
    ("Hospedaje del sitio (Cloudflare Pages)", "—", "$0 (gratis)"),
    ("Presupuesto de anuncios (opcional, Fase 2)", "Mensual", "Lo defines tú"),
], col2="Costo (USD)")
add_body(
    "El hospedaje en Cloudflare Pages es gratuito (igual que en nuestros otros proyectos). El correo "
    "profesional lo dejamos sin costo mensual usando reenvío hacia tu correo actual; solo si prefieres "
    "un buzón completo tipo Google Workspace habría una mensualidad.",
    size=10, color=GRIS,
)
add_body(
    "* Rangos de referencia ajustables. Recomendación: comenzar con el paquete Fase 0 + Fase 1.",
    size=10, color=GRIS,
)

# Paquete recomendado (banda)
add_sub("Paquete recomendado para empezar")
tbl = doc.add_table(rows=1, cols=1)
cell = tbl.rows[0].cells[0]
shade_cell(cell, GRIS_CLARO_HEX)
set_cell_text(
    cell,
    "Fase 0 + Fase 1: logo, Perfil de Empresa en Google, sistema de reseñas y sitio bilingüe de "
    "conversión. Costos de terceros mínimos: solo el dominio (~$10–$15 al año); correo y hospedaje sin "
    "costo mensual. Es el conjunto mínimo que te hace visible y convierte visitas en llamadas.",
    size=10.5,
)

# ============================================================
# 6. CRONOGRAMA
# ============================================================
add_bar_heading("6. Tiempos estimados")
add_bullet("Fase 0 (cimientos): 3 a 5 días hábiles una vez recibida tu información y fotos.", bold_prefix="")
add_bullet("Fase 1 (sitio bilingüe): 5 a 8 días hábiles adicionales.", bold_prefix="")
add_body(
    "Nota de temporada: como tu trabajo va de mayo a octubre, entre más pronto lancemos, más aprovechamos "
    "lo que resta de esta temporada y llegamos posicionados a la próxima.",
    italic=True,
)

# ============================================================
# 7. QUÉ NECESITAMOS DE TI
# ============================================================
add_bar_heading("7. Qué necesitamos de ti")
add_bullet("El nombre legal exacto de tu empresa tal como quedó registrada la LLC.")
add_bullet("El nombre de dominio que prefieres (por ejemplo, lopezsealcoating.com), para registrarlo a tu nombre.")
add_bullet("Confirmar el teléfono que se mostrará al público (actualmente 331-236-9387).")
add_bullet("El correo personal donde quieres recibir los mensajes MIENTRAS configuramos tu correo profesional nuevo.")
add_bullet("Tus fotos y videos de trabajos terminados en buena calidad.")
add_bullet("Confirmar si cuentas con seguro (para poder poner \u201cLicensed & Insured\u201d).")
add_bullet("Las ciudades/zonas donde más te interesa conseguir clientes.")

# ============================================================
# 8. CONSIDERACIONES IMPORTANTES
# ============================================================
add_bar_heading("8. Consideraciones importantes")
add_body(
    "Estos puntos son informativos y conviene confirmarlos con tu contador o con la fuente oficial; "
    "no constituyen asesoría legal:",
    italic=True, size=10, color=GRIS,
)
add_bullet("Logo: una LLC puede usar cualquier logo, siempre que sea original. Por eso rehacemos el tuyo desde cero, para que sea totalmente tuyo y sin riesgos de derechos de autor.", bold_prefix="")
add_bullet("Nombre legal: en documentos formales (cotizaciones, contratos, facturas) conviene mostrar el nombre completo, incluyendo \u201cLLC\u201d.", bold_prefix="")
add_bullet("Licencias/seguro: Illinois no tiene licencia estatal de contratista; cada ciudad o condado puede pedir su propio registro. El seguro de responsabilidad ayuda mucho, sobre todo para clientes comerciales.", bold_prefix="")
add_bullet("En el sitio solo publicaremos sellos y afirmaciones verdaderas y verificables.", bold_prefix="")

# ============================================================
# 9. SIGUIENTE PASO
# ============================================================
add_bar_heading("9. Siguiente paso")
add_body(
    "Si esta propuesta te hace sentido, el siguiente paso es aprobar el paquete recomendado (Fase 0 + Fase 1) "
    "y agendar una llamada corta para recopilar tu información y fotos. A partir de ahí arrancamos de inmediato."
)

# Bloque de aceptación
acc = doc.add_table(rows=2, cols=2)
acc.style = "Table Grid"
acc.alignment = WD_TABLE_ALIGNMENT.CENTER
set_cell_text(acc.rows[0].cells[0], "Nombre y firma (Cliente)", bold=True, size=10)
set_cell_text(acc.rows[0].cells[1], "Fecha", bold=True, size=10)
set_cell_text(acc.rows[1].cells[0], "\n\n", size=10)
set_cell_text(acc.rows[1].cells[1], "\n\n", size=10)

doc.add_paragraph()

# ---------------- Firma ----------------
add_body("Atentamente,", space_after=10)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
r = p.add_run("Antonio Aguilar")
r.bold = True
r.font.size = Pt(12.5)
r.font.color.rgb = ASFALTO
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
r = p.add_run("CEO · Aimarktech")
r.font.size = Pt(11)
r.font.color.rgb = GRIS
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(10)
r = p.add_run("https://soyaimarktech.com")
r.font.size = Pt(11)
r.bold = True
r.font.color.rgb = ASFALTO

# ============================================================
# CONTACTO (pie)
# ============================================================
tbl = doc.add_table(rows=1, cols=1)
cell = tbl.rows[0].cells[0]
shade_cell(cell, ASFALTO_HEX)
p = cell.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("AIMARKTECH")
r.bold = True
r.font.size = Pt(14)
r.font.color.rgb = AMARILLO
p2 = cell.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
p2.paragraph_format.space_after = Pt(2)
r = p2.add_run("Gracias por tu confianza. Estamos listos para hacer crecer Lopez Sealcoating.")
r.font.size = Pt(10.5)
r.font.color.rgb = BLANCO
p3 = cell.add_paragraph()
p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
p3.paragraph_format.space_after = Pt(8)
r = p3.add_run("www.soyaimarktech.com")
r.font.size = Pt(10.5)
r.bold = True
r.font.color.rgb = AMARILLO

# Guardar
out = "Propuesta_Presencia_Digital_Lopez_Sealcoating.docx"
doc.save(out)
print("OK ->", out)
